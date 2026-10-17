from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_font(run, name="微軟正黑體", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    if level == 1:
        set_font(run, size=18, bold=True, color=(15, 23, 42))
    elif level == 2:
        set_font(run, size=14, bold=True, color=(30, 58, 138))
    else:
        set_font(run, size=12, bold=True, color=(51, 65, 85))
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(bold_prefix)
    set_font(r1, size=10.5, bold=True, color=(15, 23, 42))
    r2 = p.add_run(text)
    set_font(r2, size=10.5)

def build_spec_docx():
    doc = Document()

    # 文件大標題
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("郵件智慧派工與動態輪值追蹤系統\n系統設計與開發規格書 (Zero-DB Edition)")
    set_font(title_run, size=20, bold=True, color=(180, 83, 9))

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(
        "文件版本：v2.0 (Cloudflare KV / R2 單一 JSON 儲存架構)\n"
        "更新日期：2026-10-17\n"
        "目標開發環境：Google Antigravity / Next.js 14+ (Vercel) + Cloudflare Workers"
    )
    set_font(meta_run, size=10, color=(100, 116, 139))

    # 1. 架構說明
    add_heading_styled(doc, "1. 專案概述與免資料庫架構 (Zero-DB Architecture)", level=2)
    p1 = doc.add_paragraph()
    set_font(p1.add_run(
        "本系統採用無傳統資料庫（Zero-DB）設計，所有狀態（輪值人員名冊、寄件者白名單、工單紀錄與歷程、郵件攔截日誌）"
        "統一封裝為單一結構化 JSON 檔案（system_state.json），儲存於 Cloudflare R2 或 Cloudflare KV 中，"
        "並在本地開發時自動 Fallback 至本地檔案（.data/system_state.json）。"
    ), size=10.5)

    # 2. 郵件解析與業務邏輯
    add_heading_styled(doc, "2. 核心業務與郵件解析規格 (Core Business Rules)", level=2)
    add_bullet(doc, "寄件者白名單機制：", "僅允許列於 allowed_senders 白名單（支援完整 Email 或 *@domain.com 萬用字元）之寄件者觸發新工單派發；團隊成員回信則允許更新工單狀態。")
    add_bullet(doc, "僅解析最新一封回覆：", "自動偵測並截斷 Re:/Fwd: 郵件中的歷史引文（如 On ... wrote:、在 ... 寫道：、-----Original Message-----、From:/寄件者: 及 > 引用行），僅解析最上方最新回覆內容。")
    add_bullet(doc, "解析 @執行人員（支援多人）：", "掃描主旨與最新內文中的 @人員（支援半形 @ 與全形 ＠，比對姓名與 aliases 別名陣列）。單封信可同時 @多人，系統為每位被點名者建立對應工單並累計工作量。")
    add_bullet(doc, "解析 P1~P5 優先順序（預設 P3）：", "緊隨 @人員 後方解析 P1~P5（如 @建宏 P1 @Alex P2）。P1=今天回覆、P2=明天回覆、P3=3天內回覆（未標註P時之預設值）、P4=一星期內、P5=視需求。網頁端亦支援直接下拉調整優先順序。")
    add_bullet(doc, "解析 MAIL 基本欄位：", "自動擷取 MAIL 主旨（含去除 Re:/Fwd: 之乾淨主旨）、MAIL 寄送時間、MAIL 寄送人員名稱與信箱。")
    add_bullet(doc, "工單狀態與 Email 回報更新：", "狀態包含待處理 (PENDING)、處理中 (IN_PROGRESS)、已處理 (RESOLVED)、已結案 (CLOSED)。支援於網頁切換，或由負責人員直接回信內文包含 #處理中、#已處理、#結案 自動更新狀態。")
    add_bullet(doc, "雙模式派工與休假動態平衡：", "信件有 @人員 走『指定派單』；未指定人員走『系統平衡輪值』，自動指派給在線且有效負載權重 (effectiveLoad) 最低者，並將皇冠 👑 移交給最新受派者。當人員休假銷假恢復上班時，系統自動將其 effectiveLoad 對齊為 (在線最低值 - 1)，確保銷假後優先第一順位出動輪值補單，同時防止被連續爆單。")
    add_bullet(doc, "任務轉派機制：", "網頁提供轉派按鈕與原因輸入，轉派後原處理人計數 -1，新受派人計數 +1 並接手皇冠。")

    # 3. 優先級對照表
    add_heading_styled(doc, "3. 優先級 (P1~P5) 與 SLA 到期日計算標準", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, title in enumerate(["優先級代碼", "時效定義", "自動到期時間 (dueAt) 計算規則"]):
        r = hdr[i].paragraphs[0].add_run(title)
        set_font(r, size=10, bold=True)

    priority_rows = [
        ("P1", "今天要回覆", "MAIL 收信當日 23:59:59 (紅色高優警示)"),
        ("P2", "明天回覆", "MAIL 收信日 +1 天 23:59:59 (橘色標記)"),
        ("P3 (預設)", "3天內回覆", "MAIL 收信日 +3 天 23:59:59 (未標註 P 時自動套用)"),
        ("P4", "一星期內回覆", "MAIL 收信日 +7 天 23:59:59"),
        ("P5", "視需求回覆", "null (無強制期限)")
    ]
    for code, desc, rule in priority_rows:
        row_cells = table.add_row().cells
        set_font(row_cells[0].paragraphs[0].add_run(code), size=10, bold=True)
        set_font(row_cells[1].paragraphs[0].add_run(desc), size=10)
        set_font(row_cells[2].paragraphs[0].add_run(rule), size=10)

    # 4. 給 Antigravity 的開發指令
    add_heading_styled(doc, "4. 交付給 Antigravity 的啟動開發指令", level=2)
    prompt_p = doc.add_paragraph()
    prompt_run = prompt_p.add_run(
        "請依據本規格書使用 Next.js (App Router) + TypeScript + Tailwind CSS + Lucide Icons 建立完整專案：\n"
        "1. 實作 src/types/store.ts 與 src/lib/storage.ts（支援 Cloudflare R2/KV 與本地 .data/system_state.json 自動 Fallback）。\n"
        "2. 實作 src/lib/email-parser.ts（剝離舊引文只讀最新回覆、解析 @多人 + P1~P5 預設P3、解析主旨/時間/寄件人、解析回信狀態指令）。\n"
        "3. 實作 /api/webhook/email、/api/tasks/[id]、/api/roster、/api/whitelist。\n"
        "4. 實作前端儀表板 src/app/page.tsx（含頂部皇冠舞台、輪值隊列、行內切換 P1~P5 與狀態、轉派 Modal、統計圖表、人員休假維護、白名單與郵件解析測試模擬器）。"
    )
    set_font(prompt_run, name="Consolas", size=9.5, color=(30, 41, 59))

    output_name = "郵件智慧派工與動態輪值追蹤系統_開發規格書.docx"
    doc.save(output_name)
    print(f"✅ 已成功生成標準 Word 檔案：{output_name}")

if __name__ == "__main__":
    build_spec_docx()
    